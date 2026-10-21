"""Extractores sin red. Se ejecutan en un proceso separado con límites de recursos."""
import csv
import io
import json
import os
import subprocess
import sys
import warnings
import zipfile
from pathlib import Path

import fitz
from PIL import Image, ImageOps
from docx import Document
from pptx import Presentation

from app.config import Settings

FORMATS = {".pdf": "PDF", ".png": "PNG", ".jpg": "JPEG", ".jpeg": "JPEG", ".webp": "WEBP",
           ".docx": "Word", ".pptx": "PowerPoint"}
Image.MAX_IMAGE_PIXELS = 20_000_000
warnings.simplefilter("error", Image.DecompressionBombWarning)


class MaterialError(ValueError):
    pass


def validate_file(name: str, data: bytes):
    ext = Path(name).suffix.lower()
    if ext not in FORMATS:
        raise MaterialError("Formato no admitido. Usa PDF, JPG, PNG, WebP, DOCX o PPTX.")
    if not data:
        raise MaterialError("El archivo está vacío.")
    if ext == ".pdf" and not data.startswith(b"%PDF-"):
        raise MaterialError("El contenido no corresponde a un PDF.")
    if ext in {".png", ".jpg", ".jpeg", ".webp"}:
        try:
            with Image.open(io.BytesIO(data)) as img:
                if img.format != FORMATS[ext] or img.width * img.height > Image.MAX_IMAGE_PIXELS:
                    raise MaterialError("Imagen incompatible o superior a 20 megapíxeles.")
                img.verify()
        except MaterialError:
            raise
        except Exception as exc:
            raise MaterialError("La imagen está dañada o no corresponde a su extensión.") from exc
    if ext in {".docx", ".pptx"}:
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                entries = z.infolist()
                if len(entries) > 2000 or sum(x.file_size for x in entries) > 60 * 1024 * 1024:
                    raise MaterialError("Documento comprimido demasiado grande.")
                for e in entries:
                    if e.flag_bits & 1 or e.filename.startswith("/") or ".." in Path(e.filename).parts:
                        raise MaterialError("Documento cifrado o estructura no admitida.")
                    member = e.filename.lower()
                    if "vbaproject" in member or "/embeddings/" in member or (member.endswith(".bin") and "/printersettings/" not in member):
                        raise MaterialError("No se admiten macros ni objetos incrustados.")
                expected = "word/document.xml" if ext == ".docx" else "ppt/presentation.xml"
                if expected not in z.namelist():
                    raise MaterialError("El documento no corresponde a su extensión.")
        except zipfile.BadZipFile as exc:
            raise MaterialError("El documento está dañado.") from exc
    return ext


def ocr(image: Image.Image, cfg: Settings):
    image = ImageOps.exif_transpose(image).convert("RGB")
    image.thumbnail((2800, 2800))
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    command = ["tesseract", "stdin", "stdout", "--tessdata-dir", str(cfg.tessdata_dir.resolve()),
               "-l", "spa", "--psm", "3", "-c", "tessedit_create_tsv=1"]
    result = subprocess.run(command, input=buf.getvalue(), capture_output=True, timeout=25)
    if result.returncode:
        raise MaterialError("OCR no disponible. Verifica Tesseract y el modelo de español en el servidor.")
    words, confidences = [], []
    for row in csv.DictReader(io.StringIO(result.stdout.decode("utf-8")), delimiter="\t"):
        word = (row.get("text") or "").strip()
        confidence = float(row.get("conf", "-1"))
        if word and confidence >= 0:
            words.append(word)
            confidences.append(confidence)
    text = " ".join(words)
    confidence = sum(confidences) / max(1, len(confidences))
    if len(text) < 20 or confidence < 65:
        raise MaterialError("No se lee con suficiente confianza. Sube una foto más nítida, recta y bien iluminada, o transcribe ese fragmento en un documento.")
    return text, round(confidence, 1)


def extract(path: Path, cfg: Settings):
    ext = path.suffix.lower()
    pages, notes = [], []

    def add(text, page, method="text", confidence=None):
        text = " ".join(text.split())
        if len(text) < 20:
            raise MaterialError(f"No hay texto suficiente en {'la página ' + str(page) if page else 'el documento'}. Adjunta una versión legible.")
        pages.append({"page": page, "text": text, "method": method, "confidence": confidence})
        if sum(len(p["text"]) for p in pages) > cfg.max_text_chars:
            raise MaterialError(f"Supera {cfg.max_text_chars} caracteres extraídos. Divide el material.")

    if ext == ".pdf":
        with fitz.open(path) as doc:
            if doc.is_encrypted:
                raise MaterialError("El PDF tiene contraseña. Adjunta una copia sin contraseña.")
            if len(doc) > cfg.max_pages:
                raise MaterialError(f"El PDF supera el límite de {cfg.max_pages} páginas.")
            for i, page in enumerate(doc):
                text = page.get_text()
                if len(text.strip()) >= 20:
                    add(text, i + 1)
                    if page.get_images():
                        notes.append(f"Página {i+1}: se extrajo el texto; revisa diagramas e imágenes por separado.")
                else:
                    scale = min(2, 2800 / max(page.rect.width, page.rect.height))
                    pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
                    text, conf = ocr(Image.frombytes("RGB", [pix.width, pix.height], pix.samples), cfg)
                    add(text, i + 1, "ocr", conf)
    elif ext in {".png", ".jpg", ".jpeg", ".webp"}:
        with Image.open(path) as image:
            text, conf = ocr(image, cfg)
        add(text, 1, "ocr", conf)
    elif ext == ".docx":
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        text += "\n" + "\n".join(" ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
        add(text, None)
        notes.append("Word: texto y tablas, sin paginación estable ni OCR de imágenes incrustadas.")
    elif ext == ".pptx":
        slides = Presentation(path).slides
        if len(slides) > cfg.max_pages:
            raise MaterialError(f"La presentación supera {cfg.max_pages} diapositivas.")
        for i, slide in enumerate(slides):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text)
                if shape.has_table:
                    texts.extend(" ".join(c.text for c in row.cells) for row in shape.table.rows)
            add("\n".join(texts), i + 1)
        notes.append("PowerPoint: texto y tablas; exporta diapositivas escaneadas a PDF para OCR.")
    if not pages:
        raise MaterialError("El documento no contiene páginas con texto.")
    if any(p["method"] == "ocr" for p in pages):
        notes.append("OCR automático: comprueba el texto extraído, especialmente cifras, fórmulas y escritura a mano.")
    return {"pages": pages, "notes": list(dict.fromkeys(notes))[:10]}


if __name__ == "__main__":
    import resource
    resource.setrlimit(resource.RLIMIT_AS, (1536 * 1024 * 1024, 1536 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_CPU, (60, 60))
    try:
        print(json.dumps({"ok": True, **extract(Path(sys.argv[1]), Settings())}))
    except MaterialError as exc:
        print(json.dumps({"ok": False, "error": str(exc)}))
    except Exception:
        print(json.dumps({"ok": False, "error": "No se pudo leer el documento. Puede estar dañado o usar una estructura no admitida."}))
